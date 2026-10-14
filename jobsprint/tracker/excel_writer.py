"""The structured Excel application tracker (requirement #7).

One applications.xlsx per user, columns exactly as specified. openpyxl is
used directly rather than pandas — this file is written to incrementally,
a handful of rows per session, so there's no benefit to a heavier dataframe
dependency for what's really just "load, append, save".
"""

import os
import threading
from datetime import date

from openpyxl import Workbook, load_workbook

COLUMNS = [
    "Date & Time", "Company", "Job Role", "Experience", "Salary/CTC",
    "Posted Date", "Naukri URL", "External Application URL",
    "Application Type", "Status",
]

SUMMARY_SHEET = "Daily Summary"
SUMMARY_COLUMNS = ["Date", "Applied", "Failed", "Ignored", "Skipped", "Total"]


def _session_counts(bot):
    """The per-session counters every bot keeps, keyed by summary column."""
    return {
        "Applied": getattr(bot, "applied_count", 0),
        "Failed": getattr(bot, "failed_count", 0),
        "Ignored": getattr(bot, "ignored_count", 0),
        "Skipped": getattr(bot, "skipped_count", 0),
    }


class ExcelTracker:
    def __init__(self, path: str):
        self.path = path
        self._lock = threading.Lock()
        self._ensure_workbook()

    def _ensure_workbook(self):
        if os.path.exists(self.path):
            return
        os.makedirs(os.path.dirname(self.path), exist_ok=True)
        wb = Workbook()
        ws = wb.active
        ws.title = "Applications"
        ws.append(COLUMNS)
        for col_idx, header in enumerate(COLUMNS, start=1):
            ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = max(14, len(header) + 2)
        wb.save(self.path)

    def append_row(self, record: dict):
        row = [
            record.get("datetime", ""), record.get("company", ""),
            record.get("job_role", ""), record.get("experience", ""),
            record.get("salary", ""), record.get("posted_date", ""),
            record.get("naukri_url", ""), record.get("external_url", ""),
            record.get("application_type", ""), record.get("status", ""),
        ]
        with self._lock:
            wb = load_workbook(self.path)
            ws = wb["Applications"]
            ws.append(row)
            wb.save(self.path)

    def update_status(self, naukri_url: str, status: str):
        if not naukri_url:
            return
        with self._lock:
            wb = load_workbook(self.path)
            ws = wb["Applications"]
            url_col = COLUMNS.index("Naukri URL") + 1
            status_col = COLUMNS.index("Status") + 1
            for row in ws.iter_rows(min_row=2):
                if row[url_col - 1].value == naukri_url:
                    row[status_col - 1].value = status
                    break
            wb.save(self.path)

    def add_daily_summary(self, bot, day=None):
        """Add one session's counts to today's row on the Daily Summary
        sheet — a date that already has a row gets its numbers added to,
        so running several sessions in a day still leaves one row."""
        counts = _session_counts(bot)
        if not any(counts.values()):
            return
        day = (day or date.today()).isoformat()
        with self._lock:
            wb = load_workbook(self.path)
            if SUMMARY_SHEET in wb.sheetnames:
                ws = wb[SUMMARY_SHEET]
            else:
                ws = wb.create_sheet(SUMMARY_SHEET)
                ws.append(SUMMARY_COLUMNS)
                for col_idx in range(1, len(SUMMARY_COLUMNS) + 1):
                    ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = 12

            row = next((r for r in ws.iter_rows(min_row=2) if str(r[0].value) == day), None)
            if row is None:
                ws.append([day] + [0] * (len(SUMMARY_COLUMNS) - 1))
                row = ws[ws.max_row]
            for name, value in counts.items():
                cell = row[SUMMARY_COLUMNS.index(name)]
                cell.value = int(cell.value or 0) + value
            row[SUMMARY_COLUMNS.index("Total")].value = sum(
                int(row[SUMMARY_COLUMNS.index(name)].value or 0) for name in counts
            )
            wb.save(self.path)
